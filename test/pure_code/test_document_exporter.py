"""PureCode docx 导出器单元测试（function/document_exporter.py）"""

from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from pure_code.function.document_exporter import (
    CODE_FONT_NAME,
    CODE_FONT_SIZE_PT,
    DocxExporter,
    FileBlock,
)


class TestDocxExporter:
    """docx 文档结构与排版"""

    def test_write_structure_and_order(self, tmp_path: Path):
        blocks = [
            FileBlock("src/a.py", "x = 1\n\ny = 2"),
            FileBlock("b.py", "print('中文')"),
        ]
        output = tmp_path / "out.docx"
        DocxExporter().write(output, "demo", blocks)
        texts = [p.text for p in Document(output).paragraphs]
        expected = ["demo", "src/a.py", "x = 1", "", "y = 2", "b.py", "print('中文')"]
        assert texts == expected

    def test_code_font_is_monospace(self, tmp_path: Path):
        output = tmp_path / "out.docx"
        DocxExporter().write(output, "demo", [FileBlock("a.py", "x = 1")])
        document = Document(output)
        code_paragraph = document.paragraphs[2]
        run = code_paragraph.runs[0]
        assert run.font.name == CODE_FONT_NAME
        assert run.font.size == Pt(CODE_FONT_SIZE_PT)
        assert run.font.element.rPr.rFonts.get(qn("w:eastAsia")) == CODE_FONT_NAME

    def test_write_to_invalid_path_raises(self, tmp_path: Path):
        output = tmp_path / "不存在的目录" / "out.docx"
        with pytest.raises(OSError):
            DocxExporter().write(output, "demo", [FileBlock("a.py", "x = 1")])
