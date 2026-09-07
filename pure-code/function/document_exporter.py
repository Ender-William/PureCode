"""
PureCode docx 导出适配器

以适配器模式封装 python-docx，对上仅暴露「项目名 + 有序文件块列表」的
纯粹导出接口，隔离第三方库细节，便于未来扩展其他导出格式。
"""

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

# ===== 排版常量 =====
CODE_FONT_NAME = "Consolas"
CODE_FONT_SIZE_PT = 10
TITLE_FONT_SIZE_PT = 12
PAGE_WIDTH_CM = 21.0   # A4 纵向
PAGE_HEIGHT_CM = 29.7
PAGE_MARGIN_CM = 2.54


@dataclass
class FileBlock:
    """单个文件的导出块：相对路径标题 + 已剥离注释的代码文本"""

    relative_path: str
    code: str


class DocxExporter:
    """docx 文档导出器（python-docx 适配器）"""

    def write(
        self,
        output_path: "str | Path",
        project_name: str,
        file_blocks: list[FileBlock],
    ) -> None:
        """
        生成代码 Word 文档（A4 页面，项目名标题 + 逐文件块）

        Args:
            output_path: 输出文件路径（.docx）
            project_name: 项目名称（作为文档标题）
            file_blocks: 有序文件块列表

        Raises:
            OSError: 输出路径不可写
        """
        document = Document()
        self._setup_page(document)
        document.add_heading(project_name, level=0)
        for block in file_blocks:
            self._append_file_block(document, block)
        document.save(str(output_path))

    def _setup_page(self, document: Document) -> None:
        """设置 A4 页面尺寸与页边距"""
        section = document.sections[0]
        section.page_width = Cm(PAGE_WIDTH_CM)
        section.page_height = Cm(PAGE_HEIGHT_CM)
        section.left_margin = Cm(PAGE_MARGIN_CM)
        section.right_margin = Cm(PAGE_MARGIN_CM)
        section.top_margin = Cm(PAGE_MARGIN_CM)
        section.bottom_margin = Cm(PAGE_MARGIN_CM)

    def _append_file_block(self, document: Document, block: FileBlock) -> None:
        """追加单文件块：相对路径标题段 + 逐行代码段"""
        heading = document.add_heading("", level=1)
        title_run = heading.add_run(block.relative_path)
        title_run.font.size = Pt(TITLE_FONT_SIZE_PT)
        for line in block.code.split("\n"):
            self._append_code_line(document, line)

    def _append_code_line(self, document: Document, line: str) -> None:
        """追加单行代码段：等宽字体、零段间距（空行以空段保留）"""
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(line)
        run.font.name = CODE_FONT_NAME
        run.font.size = Pt(CODE_FONT_SIZE_PT)
        # python-docx 的 font.name 仅设置西文字体，需单独设置 eastAsia
        # 保证代码中的中文字符同样使用等宽字体
        run.font.element.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT_NAME)
